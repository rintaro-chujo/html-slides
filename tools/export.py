#!/usr/bin/env python3
"""HTML deck exporter: index.html -> PDF / PNG / PPTX.

Subcommands
  pdf   [--hook A|B|C|all] [--notes] [--out PATH]
  png   [--hook X] [--dpi N]
  pptx  [--hook X] [--out PATH]
  all   [--hook X]

Global options: --html PATH (default ./index.html), --name NAME, --chrome PATH, --timeout SEC, --tmpdir DIR

  cd <発表フォルダ> && python3 ~/ghq/github.com/rintaro-chujo/html-slides/tools/export.py all

Chrome には file:// ではなく、この箱の serve.py を一時的に立てて http:// で読ませる。発表フォルダに
deck/ が無くても箱の deck/ が配信されるので、発表側にツールのコピーは要らない。出力は <発表フォルダ>/dist/。
出力名は index.html の <meta name="deck-export-name" content="MyTalk2026-slides"> で決める（無ければフォルダ名）。

Requires: Google Chrome, poppler (pdfinfo / pdftoppm; pdfseparate + pdfunite or pypdf
for trimming a trailing blank page), python-pptx, Pillow. Python 3.9+.
"""
import argparse
import html as htmlmod
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
from pathlib import Path

TOOLS_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS_DIR))
import serve  # noqa: E402  同じ箱の配信サーバ（deck/ の代替配信を含む）

CHROME_DEFAULT = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
# main() で --html から決める
DIST_DIR = None
BASENAME = None
BASE_URL = None

SLIDE_W_PX, SLIDE_H_PX = 1920, 1080
PAGE_W_PT, PAGE_H_PT = 1440.0, 810.0  # 1920x1080 css px @ 96dpi -> pt @ 72dpi
PPTX_W_EMU, PPTX_H_EMU = 12192000, 6858000  # 13.333in x 7.5in (16:9)

HOOKS = ("A", "B", "C", "all")


def log(msg):
    print("[export] " + msg, file=sys.stderr)


def warn(msg):
    print("[export] WARNING: " + msg, file=sys.stderr)


def die(msg, code=1):
    print("[export] ERROR: " + msg, file=sys.stderr)
    sys.exit(code)


def which(name):
    for cand in (name, "/opt/homebrew/bin/" + name, "/usr/local/bin/" + name):
        p = shutil.which(cand)
        if p:
            return p
    return None


# --------------------------------------------------------------------------
# Chrome
# --------------------------------------------------------------------------
class Chrome:
    """Headless Chrome driver.

    Chrome 152 on macOS finishes --print-to-pdf / --dump-dom within a few seconds but
    then never exits (hangs at shutdown). So instead of waiting for the process we wait
    for the *output* (PDF file stable / DOM tail on stdout), send SIGTERM (which makes
    Chrome flush stdout and exit rc=0), and SIGKILL the process group as a last resort.
    """

    QUIET_SEC = 2.0      # dump-dom: no new stdout bytes for this long -> assume done
    STABLE_SEC = 1.5     # print-to-pdf: file size unchanged for this long -> assume done

    def __init__(self, binary, tmpdir, timeout):
        self.binary = binary
        self.tmpdir = Path(tmpdir)
        self.timeout = timeout
        if not Path(binary).exists():
            die("Chrome not found at %s (use --chrome or $CHROME)" % binary)
        self.tmpdir.mkdir(parents=True, exist_ok=True)

    def _cmd(self, url, profile, extra_args):
        return [
            self.binary,
            "--headless=new",
            "--disable-gpu",
            "--no-first-run",
            "--no-default-browser-check",
            "--disable-extensions",
            "--disable-sync",
            "--disable-background-networking",
            "--disable-component-update",
            "--hide-scrollbars",
            "--user-data-dir=" + profile,
            "--virtual-time-budget=8000",
            "--run-all-compositor-stages-before-draw",
        ] + list(extra_args) + [url]

    @staticmethod
    def _stop(proc):
        """SIGTERM (graceful: flushes stdout), then SIGKILL the whole process group."""
        if proc.poll() is not None:
            return
        try:
            proc.terminate()
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            try:
                os.killpg(proc.pid, signal.SIGKILL)
            except OSError:
                proc.kill()
            proc.wait()

    def _popen(self, url, profile, extra_args, stdout):
        return subprocess.Popen(
            self._cmd(url, profile, extra_args),
            stdout=stdout, stderr=subprocess.DEVNULL, start_new_session=True,
        )

    def print_to_pdf(self, url, out_pdf):
        out_pdf = Path(out_pdf)
        out_pdf.parent.mkdir(parents=True, exist_ok=True)
        if out_pdf.exists():
            out_pdf.unlink()
        profile = tempfile.mkdtemp(prefix="chrome-profile-", dir=str(self.tmpdir))
        proc = self._popen(url, profile, ["--no-pdf-header-footer", "--print-to-pdf=" + str(out_pdf)],
                           subprocess.DEVNULL)
        t0 = time.time()
        size, since = -1, None
        try:
            while time.time() - t0 < self.timeout:
                if proc.poll() is not None:
                    break
                if out_pdf.exists():
                    s = out_pdf.stat().st_size
                    if s != size:
                        size, since = s, time.time()
                    elif s > 0 and time.time() - since >= self.STABLE_SEC:
                        break
                time.sleep(0.1)
        finally:
            self._stop(proc)
            shutil.rmtree(profile, ignore_errors=True)
        if not out_pdf.exists() or out_pdf.stat().st_size == 0:
            die("Chrome did not produce %s within %ss (url: %s)" % (out_pdf, self.timeout, url))
        return out_pdf

    def dump_dom(self, url):
        profile = tempfile.mkdtemp(prefix="chrome-profile-", dir=str(self.tmpdir))
        proc = self._popen(url, profile, ["--dump-dom"], subprocess.PIPE)
        buf = bytearray()
        last = [None]

        def reader():
            while True:
                chunk = proc.stdout.read(4096)
                if not chunk:
                    break
                buf.extend(chunk)
                last[0] = time.time()

        th = threading.Thread(target=reader, daemon=True)
        th.start()
        t0 = time.time()
        try:
            while time.time() - t0 < self.timeout:
                if proc.poll() is not None:
                    break
                if b"</html>" in buf:
                    break
                # Chrome block-buffers stdout: the tail arrives only after SIGTERM.
                if last[0] is not None and time.time() - last[0] > self.QUIET_SEC:
                    break
                time.sleep(0.05)
        finally:
            self._stop(proc)          # flushes the remaining bytes
            th.join(5)
            shutil.rmtree(profile, ignore_errors=True)
        if not buf:
            die("Chrome --dump-dom produced no output within %ss (url: %s)" % (self.timeout, url))
        return buf.decode("utf-8", "replace")


# --------------------------------------------------------------------------
# URL / manifest
# --------------------------------------------------------------------------
def deck_url(html_path, hook, notes=False):
    params = {"mode": "print", "hook": hook}
    if notes:
        params["notes"] = "1"
    return "%s/%s?%s" % (BASE_URL, urllib.parse.quote(Path(html_path).name), urllib.parse.urlencode(params))


def start_server(root):
    """発表フォルダを配信する serve.py のサーバを空きポートで裏に立てる。"""
    from http.server import ThreadingHTTPServer
    srv = ThreadingHTTPServer(("127.0.0.1", 0), serve.make_handler(str(root), None))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv, "http://127.0.0.1:%d" % srv.server_address[1]


def export_name(html_path):
    src = Path(html_path).read_text(encoding="utf-8")
    m = re.search(r'<meta\s+name="deck-export-name"\s+content="([^"]+)"', src)
    return m.group(1) if m else Path(html_path).resolve().parent.name


def manifest_from_dom(dom):
    m = re.search(
        r'<script[^>]*id="deck-manifest"[^>]*>(.*?)</script>', dom, re.S | re.I
    )
    if not m:
        return None
    try:
        data = json.loads(m.group(1))
    except json.JSONDecodeError as e:
        warn("deck-manifest JSON could not be parsed: %s" % e)
        return None
    return data if isinstance(data, list) else None


_TAG_RE = re.compile(r"<[^>]+>")
_BLOCK_END_RE = re.compile(r"</(p|div|li|h[1-6]|br)\s*>|<br\s*/?>", re.I)


def _html_to_text(fragment):
    txt = _BLOCK_END_RE.sub("\n", fragment)
    txt = _TAG_RE.sub("", txt)
    txt = htmlmod.unescape(txt)
    lines = [re.sub(r"\s+", " ", ln).strip() for ln in txt.splitlines()]
    return "\n".join(ln for ln in lines if ln)


def _attr(attrs, name):
    m = re.search(r'\b%s\s*=\s*(?:"([^"]*)"|\'([^\']*)\'|([^\s>]+))' % re.escape(name), attrs)
    if not m:
        return None
    return next(g for g in m.groups() if g is not None)


def manifest_from_html(html_path, hook):
    """Fallback: regex-parse <section class="slide"> blocks (no nested <section> assumed)."""
    src = Path(html_path).read_text(encoding="utf-8")
    deck = re.search(r'<div\s+id="deck"[^>]*>(.*?)</div>\s*<!--\s*発表者ビュー', src, re.S)
    body = deck.group(1) if deck else src
    out = []
    for m in re.finditer(r"<section\b([^>]*)>(.*?)</section>", body, re.S | re.I):
        attrs, inner = m.group(1), m.group(2)
        cls = _attr(attrs, "class") or ""
        if "slide" not in cls.split():
            continue
        h = _attr(attrs, "data-hook")
        if h and hook != "all" and h != hook:
            continue
        t = re.search(r"<(h1|h2)\b[^>]*>(.*?)</\1>", inner, re.S | re.I)
        title = _html_to_text(t.group(2)).replace("\n", " ") if t else ""
        n = re.search(r'<aside\s+class="notes"[^>]*>(.*?)</aside>', inner, re.S | re.I)
        notes = _html_to_text(n.group(1)) if n else ""
        out.append({
            "index": len(out) + 1,
            "id": _attr(attrs, "id"),
            "hook": h,
            "backup": _attr(attrs, "data-backup") == "1",
            "min": float(_attr(attrs, "data-min") or 0),
            "title": title,
            "notes": notes,
        })
    return out


def _squash(s):
    return re.sub(r"\s+", "", s or "")


def get_manifest(chrome, html_path, hook):
    url = deck_url(html_path, hook)
    dom = chrome.dump_dom(url)
    man = manifest_from_dom(dom)
    src = "dom"
    if man is None:
        warn("deck-manifest not found in DOM; falling back to regex parse of %s" % html_path)
        man = manifest_from_html(html_path, hook)
        src = "regex"
    else:
        # The manifest's `notes` is innerText of a display:none <aside>, which drops
        # paragraph breaks. When the static parse agrees on the same text, prefer its
        # version with line breaks (nicer in PowerPoint / notes pages).
        try:
            static = manifest_from_html(html_path, hook)
            if len(static) == len(man):
                for a, b in zip(man, static):
                    if (a.get("id") == b.get("id")) and _squash(a.get("notes")) == _squash(b.get("notes")):
                        a["notes"] = b["notes"]
        except Exception as e:  # never let the nicety break the export
            warn("static notes parse skipped: %s" % e)
    log("manifest (%s): %d visible slide(s) for hook=%s" % (src, len(man), hook))
    return man


# --------------------------------------------------------------------------
# PDF verification / trimming
# --------------------------------------------------------------------------
def pdf_info(pdf_path):
    """Return (pages, [(w_pt, h_pt), ...]) using pdfinfo; (None, []) if unavailable."""
    exe = which("pdfinfo")
    if not exe:
        return None, []
    proc = subprocess.run([exe, str(pdf_path)], capture_output=True, text=True)
    m = re.search(r"^Pages:\s+(\d+)", proc.stdout, re.M)
    pages = int(m.group(1)) if m else None
    sizes = []
    if pages:
        proc = subprocess.run([exe, "-f", "1", "-l", str(pages), str(pdf_path)],
                              capture_output=True, text=True)
        for mm in re.finditer(r"^Page\s+\d+ size:\s+([\d.]+) x ([\d.]+) pts", proc.stdout, re.M):
            sizes.append((float(mm.group(1)), float(mm.group(2))))
    return pages, sizes


def page_is_blank(pdf_path, page_no, tmpdir):
    """Render one page at low dpi and check whether it is uniformly (near-)white."""
    exe = which("pdftoppm")
    if not exe:
        return False
    try:
        from PIL import Image, ImageChops
    except ImportError:
        return False
    prefix = Path(tmpdir) / ("blankcheck-%d" % page_no)
    subprocess.run([exe, "-png", "-r", "20", "-f", str(page_no), "-l", str(page_no),
                    str(pdf_path), str(prefix)], capture_output=True)
    cands = sorted(Path(tmpdir).glob("blankcheck-%d*.png" % page_no))
    if not cands:
        return False
    img = Image.open(cands[0]).convert("L")
    bg = Image.new("L", img.size, 255)
    bbox = ImageChops.difference(img, bg).point(lambda v: 255 if v > 8 else 0).getbbox()
    for c in cands:
        c.unlink()
    return bbox is None


def drop_trailing_pages(pdf_path, keep, tmpdir):
    """Keep only the first `keep` pages. Uses pypdf if present, else pdfseparate+pdfunite."""
    pdf_path = Path(pdf_path)
    try:
        import pypdf  # optional
        r = pypdf.PdfReader(str(pdf_path))
        w = pypdf.PdfWriter()
        for i in range(keep):
            w.add_page(r.pages[i])
        tmp = pdf_path.with_suffix(".trim.pdf")
        with open(tmp, "wb") as f:
            w.write(f)
        tmp.replace(pdf_path)
        return True
    except ImportError:
        pass
    sep, uni = which("pdfseparate"), which("pdfunite")
    if not (sep and uni):
        return False
    work = Path(tempfile.mkdtemp(prefix="trim-", dir=str(tmpdir)))
    subprocess.run([sep, "-f", "1", "-l", str(keep), str(pdf_path), str(work / "p-%d.pdf")],
                   check=True, capture_output=True)
    parts = [str(work / ("p-%d.pdf" % i)) for i in range(1, keep + 1)]
    tmp = pdf_path.with_suffix(".trim.pdf")
    subprocess.run([uni] + parts + [str(tmp)], check=True, capture_output=True)
    tmp.replace(pdf_path)
    shutil.rmtree(work, ignore_errors=True)
    return True


def verify_pdf(pdf_path, expected_pages, tmpdir, label="PDF"):
    pages, sizes = pdf_info(pdf_path)
    if pages is None:
        warn("pdfinfo not available; skipping page verification for %s" % pdf_path)
        return
    log("%s: %d page(s) -> %s" % (label, pages, pdf_path))
    # trailing blank page(s) that Chrome sometimes appends after the last break-after:page
    if expected_pages and pages > expected_pages:
        extra = list(range(expected_pages + 1, pages + 1))
        if all(page_is_blank(pdf_path, p, tmpdir) for p in extra):
            log("trailing blank page(s) %s detected; trimming to %d page(s)" % (extra, expected_pages))
            if drop_trailing_pages(pdf_path, expected_pages, tmpdir):
                pages, sizes = pdf_info(pdf_path)
            else:
                warn("could not trim trailing pages (need pypdf or pdfseparate+pdfunite)")
        else:
            warn("%s has %d pages but %d slides are visible; extra pages are NOT blank "
                 "(overflowing content?)" % (label, pages, expected_pages))
    if expected_pages and pages != expected_pages:
        warn("%s page count %d != visible slide count %d" % (label, pages, expected_pages))
    bad = [(i + 1, s) for i, s in enumerate(sizes)
           if abs(s[0] - PAGE_W_PT) > 0.6 or abs(s[1] - PAGE_H_PT) > 0.6]
    if bad:
        warn("%s: %d page(s) not %gx%g pt: %s" % (label, len(bad), PAGE_W_PT, PAGE_H_PT, bad[:5]))
    elif sizes:
        log("%s: all pages are %gx%g pt (1920x1080 px)" % (label, PAGE_W_PT, PAGE_H_PT))


# --------------------------------------------------------------------------
# Commands
# --------------------------------------------------------------------------
def variant_tag(html_path):
    """index.html -> ""、index-abstract.html -> "-abstract"（出力名に付ける版の識別子）"""
    stem = Path(html_path).stem
    if stem == "index":
        return ""
    return "-" + (stem[len("index-"):] if stem.startswith("index-") else stem)


def default_pdf_path(args, notes=False):
    return DIST_DIR / ("%s%s%s.pdf" % (BASENAME, variant_tag(args.html), "-notes" if notes else ""))


def cmd_pdf(args, chrome, notes=None, out=None):
    notes = args.notes if notes is None else notes
    out = Path(out or args.out or default_pdf_path(args, notes))
    manifest = get_manifest(chrome, args.html, args.hook)
    url = deck_url(args.html, args.hook, notes)
    log("printing %s" % url)
    chrome.print_to_pdf(url, out)
    verify_pdf(out, len(manifest), chrome.tmpdir, label="notes PDF" if notes else "slides PDF")
    return out, manifest


def ensure_slides_pdf(args, chrome):
    pdf = default_pdf_path(args)
    if getattr(args, "reuse_pdf", False) and pdf.exists():
        log("reusing existing %s" % pdf)
        return pdf, get_manifest(chrome, args.html, args.hook)
    return cmd_pdf(args, chrome, notes=False, out=pdf)


def cmd_png(args, chrome, pdf=None, manifest=None):
    if pdf is None:
        pdf, manifest = ensure_slides_pdf(args, chrome)
    exe = which("pdftoppm")
    if not exe:
        die("pdftoppm not found (brew install poppler)")
    outdir = DIST_DIR / ("png%s" % variant_tag(args.html))
    if outdir.exists():
        shutil.rmtree(outdir)
    outdir.mkdir(parents=True)
    subprocess.run([exe, "-png", "-r", str(args.dpi), str(pdf), str(outdir / "slide")],
                   check=True, capture_output=True)
    files = sorted(outdir.glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
    final = []
    for i, f in enumerate(files, 1):
        dst = outdir / ("slide-%03d.png" % i)
        f.rename(dst)
        final.append(dst)
    # verify size
    try:
        from PIL import Image
        expect = (round(PAGE_W_PT / 72 * args.dpi), round(PAGE_H_PT / 72 * args.dpi))
        sizes = {Image.open(p).size for p in final}
        if sizes == {expect}:
            log("png: %d file(s), all %dx%d -> %s" % (len(final), expect[0], expect[1], outdir))
        else:
            warn("png sizes %s (expected %s)" % (sorted(sizes), expect))
    except ImportError:
        log("png: %d file(s) -> %s (Pillow missing, size unverified)" % (len(final), outdir))
    if manifest and len(final) != len(manifest):
        warn("png count %d != visible slide count %d" % (len(final), len(manifest)))
    return final, manifest


def cmd_pptx(args, chrome, pngs=None, manifest=None):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        die("python-pptx not installed (pip install python-pptx)")
    if pngs is None:
        pngs, manifest = cmd_png(args, chrome)
    if manifest is None:
        manifest = get_manifest(chrome, args.html, args.hook)
    out = Path(args.out or (DIST_DIR / ("%s%s.pptx" % (BASENAME, variant_tag(args.html)))))
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Emu(PPTX_W_EMU)
    prs.slide_height = Emu(PPTX_H_EMU)
    blank = prs.slide_layouts[6]
    for i, png in enumerate(pngs):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        m = manifest[i] if i < len(manifest) else {}
        meta = []
        if m.get("hook"):
            meta.append("つかみ案 %s" % m["hook"])
        if m.get("backup"):
            meta.append("付録")
        if m.get("min"):
            meta.append("予定 %g分" % m["min"])
        notes = (m.get("notes") or "").strip()
        head = " ・ ".join(meta)
        text = (head + "\n\n" if head else "") + notes
        s.notes_slide.notes_text_frame.text = text
        # keep the shape name readable in the PowerPoint selection pane
        s.shapes[0].name = "slide-%03d %s" % (i + 1, (m.get("title") or "")[:40])
    prs.save(str(out))
    log("pptx: %d slide(s) -> %s" % (len(prs.slides), out))
    return out


def cmd_all(args, chrome):
    pdf, manifest = cmd_pdf(args, chrome, notes=False, out=default_pdf_path(args))
    cmd_pdf(args, chrome, notes=True, out=default_pdf_path(args, notes=True))
    pngs, _ = cmd_png(args, chrome, pdf=pdf, manifest=manifest)
    args.out = None
    cmd_pptx(args, chrome, pngs=pngs, manifest=manifest)


# --------------------------------------------------------------------------
def build_parser():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--html", default="index.html", help="deck HTML (default: ./index.html)")
    p.add_argument("--name", help="出力ファイル名の頭（既定: <meta name=deck-export-name> かフォルダ名）")
    p.add_argument("--chrome", default=os.environ.get("CHROME", CHROME_DEFAULT))
    p.add_argument("--timeout", type=float, default=120, help="seconds per Chrome run (default 120)")
    p.add_argument("--tmpdir", default=None, help="scratch dir for Chrome profiles (default dist/.tmp)")
    sub = p.add_subparsers(dest="cmd", required=True)

    def common(sp):
        sp.add_argument("--hook", default="A", choices=HOOKS)

    sp = sub.add_parser("pdf", help="slides PDF (or notes pages with --notes)")
    common(sp)
    sp.add_argument("--notes", action="store_true")
    sp.add_argument("--out")

    sp = sub.add_parser("png", help="one PNG per slide via pdftoppm")
    common(sp)
    sp.add_argument("--dpi", type=int, default=96)
    sp.add_argument("--reuse-pdf", action="store_true", help="do not re-print if dist PDF exists")

    sp = sub.add_parser("pptx", help="16:9 PowerPoint (PNG per slide + speaker notes)")
    common(sp)
    sp.add_argument("--dpi", type=int, default=96)
    sp.add_argument("--out")
    sp.add_argument("--reuse-pdf", action="store_true")

    sp = sub.add_parser("all", help="pdf + notes pdf + png + pptx")
    common(sp)
    sp.add_argument("--dpi", type=int, default=96)
    return p


def main(argv=None):
    global DIST_DIR, BASENAME, BASE_URL
    args = build_parser().parse_args(argv)
    if not Path(args.html).exists():
        die("HTML not found: %s" % args.html)
    args.html = str(Path(args.html).resolve())
    deck_dir = Path(args.html).parent
    DIST_DIR = deck_dir / "dist"
    BASENAME = args.name or export_name(args.html)
    srv, BASE_URL = start_server(deck_dir)
    log("serving %s at %s" % (deck_dir, BASE_URL))
    args.notes = getattr(args, "notes", False)
    args.out = getattr(args, "out", None)
    args.dpi = getattr(args, "dpi", 96)
    tmpdir = Path(args.tmpdir) if args.tmpdir else DIST_DIR / ".tmp"
    chrome = Chrome(args.chrome, tmpdir, args.timeout)
    try:
        {"pdf": cmd_pdf, "png": cmd_png, "pptx": cmd_pptx, "all": cmd_all}[args.cmd](args, chrome)
    finally:
        srv.shutdown()
        if not args.tmpdir:
            shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    main()
